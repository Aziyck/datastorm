import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
import joblib, os, sklearn
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def calculeaza_pnl(input_file):
    df = pd.read_excel(input_file)

    for col in ["PnL_Linear", "PnL_Ridge", "PnL_Lasso"]:
        df[col] = None

    def prepare_numeric_row(row, cols):
        sub = pd.DataFrame([row])[cols].copy()
        sub = sub.apply(pd.to_numeric, errors="coerce").fillna(0)
        return sub

    for i, row in df.iterrows():
        scenariu = str(row["Scenariu"]).strip()
        paths = [f"models/{scenariu}_Linear.pkl", f"models/{scenariu}_Ridge.pkl", f"models/{scenariu}_Lasso.pkl"]
        if not all(os.path.exists(p) for p in paths):
            continue

        lin = joblib.load(paths[0])
        ridge = joblib.load(paths[1])
        lasso = joblib.load(paths[2])

        Xl = prepare_numeric_row(row, lin["X_columns"])
        Xr = prepare_numeric_row(row, ridge["X_columns"])
        Xs = prepare_numeric_row(row, lasso["X_columns"])

        df.at[i, "PnL_Linear"] = round(float(lin["model"].predict(Xl)[0]),2)
        df.at[i, "PnL_Ridge"] = round(float(ridge["model"].predict(Xr)[0]),2)
        df.at[i, "PnL_Lasso"] = round(float(lasso["model"].predict(Xs)[0]),2)

    output = input_file.replace(".xlsx", "_Rezultate.xlsx")
    df.to_excel(output, index=False)

    wb = load_workbook(output)
    ws = wb.active
    fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    for row in ws.iter_rows(min_row=2):
        for col in range(len(df.columns)-3, len(df.columns)):
            cell = row[col]
            if cell.value:
                cell.fill = fill
    wb.save(output)

    messagebox.showinfo("Succes", f"Fișier generat:\n{output}")

def alege_fisier():
    f = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
    if f:
        calculeaza_pnl(f)

root = tk.Tk()
root.title("Generare PnL - Modele")
root.geometry("400x200")

label = tk.Label(root, text="Alege fișierul Excel cu scenarii", font=("Arial", 12))
label.pack(pady=20)

button = tk.Button(root, text="Deschide fișier", command=alege_fisier, bg="#4CAF50", fg="white", font=("Arial", 11))
button.pack(pady=10)

root.mainloop()
